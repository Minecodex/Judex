"""Run inside the pinned analysis image with read-only fixture mounts."""
import hashlib
import importlib.util
import json
import pathlib
import tempfile
import unittest
import zipfile

spec = importlib.util.spec_from_file_location("parser", "/opt/judex/analyze_material.py")
parser = importlib.util.module_from_spec(spec)
spec.loader.exec_module(parser)


class OriginalParserTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        from pypdf import PdfWriter
        from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
        from PIL import Image
        cls.directory = tempfile.TemporaryDirectory()
        cls.root = pathlib.Path(cls.directory.name)
        document = Document()
        document.add_paragraph("DOCX_SOURCE_UNIQUE")
        document.add_table(rows=2, cols=2).cell(1, 1).text = "TABLE_SOURCE"
        document.save(cls.root / "acceptance.docx")
        deck = Presentation()
        for text in ("FIRST_SLIDE", "PPTX_SECOND_SLIDE_UNIQUE"):
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = text
        deck.save(cls.root / "acceptance.pptx")
        book = Workbook()
        book.active["A1"] = "FIRST_SHEET"
        book.create_sheet("Second")["B2"] = "XLSX_SECOND_SHEET_UNIQUE"
        book.save(cls.root / "acceptance.xlsx")
        writer = PdfWriter()
        for number in (1, 2):
            page = writer.add_blank_page(width=600, height=360)
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 30 300 Td (PDF_SOURCE_PAGE_{number}) Tj ET".encode())
            page[NameObject("/Contents")] = writer._add_object(stream)
        writer.write(cls.root / "acceptance.pdf")
        Image.new("RGB", (600, 360), "white").save(cls.root / "acceptance.png")
        (cls.root / "corrupt.docx").write_bytes(b"not an OOXML file")

    @classmethod
    def tearDownClass(cls):
        cls.directory.cleanup()

    def parse(self, filename, mime, offset=0, limit=100):
        file = self.root / filename
        sha = hashlib.sha256(file.read_bytes()).hexdigest()
        return parser.analyze(file, "controlled-fixed-version", sha, mime, offset, limit)

    def test_originals_and_source_locations(self):
        cases = [("acceptance.docx", "wordprocessingml.document", "DOCX_SOURCE_UNIQUE", {"paragraph", "tableRow"}),
                 ("acceptance.pptx", "presentationml.presentation", "PPTX_SECOND_SLIDE_UNIQUE", {"slideShape"}),
                 ("acceptance.xlsx", "spreadsheetml.sheet", "XLSX_SECOND_SHEET_UNIQUE", {"cell"})]
        for file, suffix, proof, location in cases:
            with self.subTest(file=file):
                result = self.parse(file, "application/vnd.openxmlformats-officedocument." + suffix)
                self.assertIn(proof, json.dumps(result, ensure_ascii=False))
                self.assertEqual({r["location"]["kind"] for r in result["records"]}, location)
                self.assertEqual(result["parserVersion"], "judex-original-3")
                self.assertTrue(result["limitations"])
        pdf = self.parse("acceptance.pdf", "application/pdf")
        self.assertEqual([r["location"]["page"] for r in pdf["records"]], [1, 2])
        self.assertIn("PDF_SOURCE_PAGE_2", pdf["records"][1]["text"])
        image = self.parse("acceptance.png", "image/png")
        self.assertEqual((image["records"][0]["width"], image["records"][0]["height"]), (600, 360))
        self.assertIn("not image understanding", image["limitations"][0])

    def test_pagination_retains_second_page(self):
        first = self.parse("acceptance.pdf", "application/pdf", limit=1)
        self.assertTrue(first["truncated"])
        second = self.parse("acceptance.pdf", "application/pdf", offset=first["coverage"]["nextOffset"], limit=1)
        self.assertEqual(second["records"][0]["location"]["page"], 2)
        self.assertFalse(second["coverage"]["hasMore"])

    def test_projected_original_without_extension(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            file = pathlib.Path(directory) / "content"
            book = Workbook()
            book.active.title = "Projected"
            book.active["A1"] = 12
            book.active["B1"] = "=A1*2"
            book.create_sheet("Second")["C3"] = "FIXED_ORIGINAL"
            book.save(file)
            sha = hashlib.sha256(file.read_bytes()).hexdigest()
            mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            result = parser.analyze(file, "projected-version", sha, mime, 0, 100)
            self.assertEqual(result["sha256"], sha)
            self.assertEqual(result["records"][2]["location"],
                             {"kind": "cell", "sheet": "Second", "cell": "C3"})
            self.assertEqual(result["records"][2]["value"], "FIXED_ORIGINAL")
            formula = result["records"][1]
            self.assertEqual(formula["formula"], "=A1*2")
            self.assertIsNone(formula["cachedValue"])
            self.assertFalse(formula["formulaRecalculated"])

    def test_formula_does_not_claim_recalculation(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as directory:
            file = pathlib.Path(directory) / "formula.xlsx"
            book = Workbook()
            book.active["A1"] = 12
            book.active["B1"] = "=A1*2"
            book.save(file)
            sha = hashlib.sha256(file.read_bytes()).hexdigest()
            result = parser.analyze(file, "formula-version", sha, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", 0, 50)
            formula = next(r for r in result["records"] if "formula" in r)
            self.assertEqual(formula["formula"], "=A1*2")
            self.assertIsNone(formula["cachedValue"])
            self.assertFalse(formula["formulaRecalculated"])

    def test_source_text_and_archive_pagination(self):
        with tempfile.TemporaryDirectory() as directory:
            root = pathlib.Path(directory)
            text = root / "requirements.md"
            text.write_text("需求第一行\nSECOND_LINE_PROOF\n", encoding="utf-8")
            sha = hashlib.sha256(text.read_bytes()).hexdigest()
            second = parser.analyze(text, "fixed", sha, "text/markdown", 1, 1)
            self.assertEqual(second["records"][0]["location"]["line"], 2)
            self.assertEqual(second["records"][0]["text"], "SECOND_LINE_PROOF")
            archive = root / "source.zip"
            with zipfile.ZipFile(archive, "w") as zipped:
                zipped.writestr("backend/server.py", "raise RuntimeError('must not execute')\nSOURCE_SECOND_LINE\n")
                zipped.writestr("web/index.html", "<p>FRONTEND_SOURCE</p>\n")
                zipped.writestr("binary.dat", b"\x00\xff")
            sha = hashlib.sha256(archive.read_bytes()).hexdigest()
            first = parser.analyze(archive, "source-fixed", sha, "application/zip", 0, 2)
            self.assertTrue(first["coverage"]["hasMore"])
            self.assertEqual(first["records"][1]["location"], {"kind": "archiveEntry", "entry": "backend/server.py", "line": 2})
            rest = parser.analyze(archive, "source-fixed", sha, "application/zip", 2, 20)
            self.assertEqual(rest["records"][0]["location"]["entry"], "web/index.html")
            self.assertTrue(rest["records"][1]["notParsed"])
            self.assertFalse(rest["coverage"]["hasMore"])

    def test_unsafe_archive_and_expansion_rejected(self):
        with tempfile.TemporaryDirectory() as directory:
            file = pathlib.Path(directory) / "unsafe.zip"
            for name in ("../outside.py", "/absolute.py", "C:/outside.py"):
                with zipfile.ZipFile(file, "w") as archive:
                    archive.writestr(name, "forbidden")
                sha = hashlib.sha256(file.read_bytes()).hexdigest()
                with self.assertRaisesRegex(ValueError, "unsafe ZIP entry"):
                    parser.analyze(file, "unsafe", sha, "application/zip", 0, 20)
            with zipfile.ZipFile(file, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("large.txt", b"x" * 1024)
            sha = hashlib.sha256(file.read_bytes()).hexdigest()
            original = parser.MAX_EXPANDED_BYTES
            try:
                parser.MAX_EXPANDED_BYTES = 512
                with self.assertRaisesRegex(ValueError, "expansion exceeds limit"):
                    parser.analyze(file, "large", sha, "application/zip", 0, 20)
            finally:
                parser.MAX_EXPANDED_BYTES = original

    def test_hash_and_corruption_rejected(self):
        file = self.root / "acceptance.docx"
        with self.assertRaisesRegex(ValueError, "SHA256 mismatch"):
            parser.analyze(file, "original", "0" * 64, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", 0, 50)
        with self.assertRaises(Exception):
            self.parse("corrupt.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")


if __name__ == "__main__":
    unittest.main(verbosity=2)
