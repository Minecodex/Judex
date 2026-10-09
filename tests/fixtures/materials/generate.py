from pathlib import Path
import json
import zipfile
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw

root = Path(__file__).parent
document = Document()
document.add_heading('共享资料真实文档', 0)
document.add_paragraph('这份材料用于共享资料体验验收，保留真实任务及文件版本。')
document.add_paragraph('DOCX_SOURCE_UNIQUE')
document.save(root / 'acceptance.docx')
workbook = Workbook()
sheet = workbook.active
sheet.title = '开发清单'
sheet.append(['工作项', '负责人', '进度'])
sheet.append(['文件卡片', '真实上传人', '完成'])
second = workbook.create_sheet('验收清单')
second.append(['验收项目', '说明'])
second.append(['XLSX_SECOND_SHEET_UNIQUE', '第二张工作表必须保留'])
for worksheet in workbook.worksheets:
    worksheet.column_dimensions['A'].width = 40
    worksheet.column_dimensions['B'].width = 36
    worksheet.column_dimensions['C'].width = 16
workbook.save(root / 'acceptance.xlsx')
slides = Presentation()
for title in ['共享资料评审', 'PPTX_SECOND_SLIDE_UNIQUE']:
    slide = slides.slides.add_slide(slides.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = '用于真实内容预览与交互验收。'
slides.save(root / 'acceptance.pptx')
pdf = canvas.Canvas(str(root / 'acceptance.pdf'))
for page in range(1, 3):
    pdf.drawString(72, 740, 'PDF_SOURCE_PAGE_' + str(page))
    pdf.drawString(72, 712, 'Original evidence for the material preview acceptance.')
    pdf.showPage()
pdf.save()
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
pdfmetrics.registerFont(UnicodeCIDFont('STSong-Light'))
chinese = canvas.Canvas(str(root / 'chinese-cid.pdf'))
chinese.setFont('STSong-Light', 16)
chinese.drawString(72, 740, '共享资料中文字体与固定版本验证')
chinese.showPage()
chinese.save()
image = Image.new('RGB', (600, 360), '#e6eee6')
draw = ImageDraw.Draw(image)
draw.rectangle((36, 36, 564, 324), fill='white', outline='#2d6651', width=3)
draw.text((70, 100), 'REAL SHARED MATERIAL', fill='#23312b')
draw.text((70, 160), 'PNG_SOURCE_UNIQUE', fill='#2d6651')
image.save(root / 'acceptance.png')
(root / 'acceptance.md').write_text('# 真实资料\n\n用于任务进度与文件版本核对。\n\n## 原始内容\n\n```text\nMARKDOWN_SOURCE_UNIQUE\n```\n', encoding='utf-8')
(root / 'acceptance.json').write_text(json.dumps({'purpose': '真实任务进展', 'proof': 'JSON_SOURCE_UNIQUE'}, ensure_ascii=False, indent=2), encoding='utf-8')
with zipfile.ZipFile(root / 'acceptance.zip', 'w') as archive:
    archive.writestr('README.md', 'ZIP_SOURCE_UNIQUE')
    archive.writestr('nested/manifest.json', '{"version": 1}')
with zipfile.ZipFile(root / 'unsafe.zip', 'w') as archive:
    archive.writestr('../escape.txt', 'must be rejected')
(root / 'corrupt.docx').write_bytes(b'not a real office document')
