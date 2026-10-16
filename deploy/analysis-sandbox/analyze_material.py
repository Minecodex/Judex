"""Read a hash-verified original in a run sandbox; no network or writes."""
import argparse
import hashlib
import itertools
import json
import pathlib
import stat
import sys
import uuid
import zipfile

PARSER_VERSION = "judex-original-3"
MAX_BYTES = 32 * 1024 * 1024
MAX_EXPANDED_BYTES = 64 * 1024 * 1024
TEXT_SUFFIXES = {".py", ".js", ".mjs", ".cjs", ".ts", ".tsx", ".html", ".css", ".md", ".txt", ".json", ".csv", ".yaml", ".yml", ".toml", ".go", ".rs", ".sql", ".sh"}


def checked_archive(file):
    archive = zipfile.ZipFile(file)
    entries = archive.infolist()
    if len(entries) > 10000 or sum(e.file_size for e in entries) > MAX_EXPANDED_BYTES:
        archive.close()
        raise ValueError("ZIP expansion exceeds limit")
    for entry in entries:
        name = entry.filename.replace("\\", "/")
        parts = pathlib.PurePosixPath(name)
        mode = entry.external_attr >> 16
        if parts.is_absolute() or ".." in parts.parts or ":" in name or stat.S_ISLNK(mode):
            archive.close()
            raise ValueError("unsafe ZIP entry")
    return archive


def records(file, mime, limitations):
    if mime.startswith("text/") or mime == "application/json":
        limitations.append("UTF-8 source text only; markup and code are not rendered or executed.")
        with file.open(encoding="utf-8-sig") as source:
            for line, text in enumerate(source, 1):
                yield {"location": {"kind": "line", "line": line}, "text": text.rstrip("\r\n")}
    elif mime in ("application/zip", "application/x-zip-compressed"):
        limitations.append("Bounded UTF-8 source entries only; binary entries are listed, nested archives are not expanded, and no code is executed.")
        with checked_archive(file) as archive:
            for entry in archive.infolist():
                if entry.is_dir():
                    continue
                location = {"kind": "archiveEntry", "entry": entry.filename}
                if pathlib.PurePosixPath(entry.filename).suffix.lower() not in TEXT_SUFFIXES or entry.file_size > 1024 * 1024:
                    yield {"location": location, "size": entry.file_size, "notParsed": True}
                    continue
                try:
                    text = archive.read(entry).decode("utf-8-sig")
                except UnicodeDecodeError:
                    yield {"location": location, "size": entry.file_size, "notParsed": True, "reason": "not UTF-8"}
                    continue
                for line, content in enumerate(text.splitlines(), 1):
                    yield {"location": {**location, "line": line}, "text": content}
    elif mime.endswith("wordprocessingml.document"):
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(file)
        limitations.append("Text/table extraction does not verify layout, images or tracked-change rendering.")
        for index, block in enumerate(document.element.body):
            if block.tag.endswith("}p"):
                paragraph = Paragraph(block, document)
                yield {"location": {"kind": "paragraph", "block": index + 1}, "text": paragraph.text, "style": paragraph.style.name}
            elif block.tag.endswith("}tbl"):
                for row_index, row in enumerate(Table(block, document).rows):
                    yield {"location": {"kind": "tableRow", "block": index + 1, "row": row_index + 1}, "cells": [cell.text for cell in row.cells]}
    elif mime.endswith("presentationml.presentation"):
        from pptx import Presentation
        limitations.append("Text/table extraction does not verify slide geometry, artwork or chart rendering.")
        for page, slide in enumerate(Presentation(file).slides, 1):
            for index, shape in enumerate(slide.shapes, 1):
                location = {"kind": "slideShape", "page": page, "shape": index}
                if shape.has_text_frame:
                    yield {"location": location, "text": shape.text}
                elif shape.has_table:
                    for row_index, row in enumerate(shape.table.rows, 1):
                        yield {"location": {**location, "row": row_index}, "cells": [cell.text for cell in row.cells]}
                else:
                    yield {"location": location, "unsupportedVisual": True, "shapeName": shape.name}
    elif mime.endswith("spreadsheetml.sheet"):
        from contextlib import ExitStack
        from openpyxl import load_workbook
        limitations.append("Formula cached values may be absent or stale. This parser does not recalculate formulas.")
        # Projected originals are named "content". MIME and verified bytes,
        # rather than a storage filename extension, select the parser.
        with ExitStack() as stack:
            formula_source = stack.enter_context(file.open("rb"))
            cache_source = stack.enter_context(file.open("rb"))
            formulas = load_workbook(formula_source, read_only=True, data_only=False, keep_links=False)
            stack.callback(formulas.close)
            cached = load_workbook(cache_source, read_only=True, data_only=True, keep_links=False)
            stack.callback(cached.close)
            for sheet in formulas:
                if sheet.max_row > 100000 or sheet.max_column > 512:
                    raise ValueError("sheet bounds exceed parser limit")
                for row, cached_row in zip(sheet.iter_rows(), cached[sheet.title].iter_rows()):
                    for cell, cache_cell in zip(row, cached_row):
                        if cell.value is None:
                            continue
                        item = {"location": {"kind": "cell", "sheet": sheet.title, "cell": cell.coordinate}}
                        if cell.data_type == "f":
                            item.update(formula=cell.value, cachedValue=cache_cell.value, formulaRecalculated=False)
                        else:
                            item["value"] = cell.value
                        yield item
    elif mime == "application/pdf":
        from pypdf import PdfReader
        limitations.append("Empty/scanned pages need OCR or visual review. Text extraction does not verify layout.")
        reader = PdfReader(file)
        if reader.is_encrypted:
            raise ValueError("encrypted PDF unsupported")
        for page, source in enumerate(reader.pages, 1):
            content = source.extract_text() or ""
            yield {"location": {"kind": "page", "page": page}, "text": content, "needsOCRorVision": not bool(content.strip())}
    elif mime in ("image/png", "image/jpeg", "image/webp"):
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 25000000
        with Image.open(file) as source:
            source.verify()
        with Image.open(file) as source:
            yield {"location": {"kind": "image"}, "width": source.width, "height": source.height, "mode": source.mode, "format": source.format}
        limitations.append("Image metadata is not image understanding. Use an admitted vision capability for semantic review.")
    else:
        raise ValueError("unsupported original MIME")


def analyze(file, version, sha, mime, offset, limit):
    if file.is_symlink() or not file.is_file() or file.stat().st_size > MAX_BYTES:
        raise ValueError("original is not a regular file within size limit")
    digest = hashlib.sha256()
    with file.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    if digest.hexdigest() != sha:
        raise ValueError("original SHA256 mismatch")
    if mime.startswith("application/vnd.openxmlformats-officedocument."):
        with checked_archive(file):
            pass
    limitations = []
    iterator = records(file, mime, limitations)
    try:
        selected = list(itertools.islice(iterator, offset, offset + limit + 1))
    finally:
        iterator.close()
    more = len(selected) > limit
    selected = selected[:limit]
    clipped = False
    for item in selected:
        for key in ("text", "cells", "value", "formula"):
            value = item.get(key)
            if isinstance(value, str) and len(value) > 4000:
                item[key] = value[:4000]
                item["contentTruncated"] = True
                clipped = True
            elif isinstance(value, list):
                if any(len(str(v)) > 2000 for v in value) or len(value) > 64:
                    item[key] = [str(v)[:2000] for v in value[:64]]
                    item["contentTruncated"] = True
                    clipped = True
    return {"versionId": version, "sha256": sha, "parserVersion": PARSER_VERSION, "mime": mime,
            "coverage": {"offset": offset, "count": len(selected), "hasMore": more, "nextOffset": offset + len(selected) if more else None},
            "truncated": more or clipped, "records": selected, "limitations": limitations}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--version", required=True, type=uuid.UUID)
    parser.add_argument("--sha", required=True)
    parser.add_argument("--mime", required=True)
    parser.add_argument("--offset", type=int, default=0)
    parser.add_argument("--limit", type=int, default=50)
    args = parser.parse_args()
    if not 0 <= args.offset <= 100000 or not 1 <= args.limit <= 100:
        raise ValueError("pagination bounds invalid")
    file = pathlib.Path("/workspace/project/versions") / str(args.version) / "content"
    result = analyze(file, str(args.version), args.sha, args.mime, args.offset, args.limit)
    raw = json.dumps(result, ensure_ascii=False, default=str)
    if len(raw.encode()) > 256 * 1024:
        raise ValueError("output exceeds limit; use a smaller page")
    print(raw)


if __name__ == "__main__":
    try:
        main()
    except Exception as error:
        print(json.dumps({"error": str(error), "parserVersion": PARSER_VERSION}), file=sys.stderr)
        sys.exit(1)
